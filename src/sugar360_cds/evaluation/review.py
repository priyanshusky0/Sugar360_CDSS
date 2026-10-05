"""The review workbook: one Excel file per approach, e.g. Approach_A.xlsx (layout: docs/decisions.md D16).

Sheets:
  How to read     what each column group means, how to fill the review columns,
                  and one real row explained
  Rows            the decision table, WHO-style: a coloured group header row
                  IF | THEN | SOURCE | AUTO CHECKS | REVIEW above the column names;
                  row_id and guideline first; header rows and first two columns
                  frozen; IF columns unused by every row are hidden
  Coverage        per guideline: pages read and never read, calls, rows, skipped items
  Conflicts       only the two automatic types (different threshold; recommend vs avoid)
  Compare by drug one block per drug class and decision, each guideline side by side; no verdict
  For dose table  dose / titration / reduce-dose rows, same layout as Rows
  Skipped         items the model did not turn into rows, with its reason category

Files are labelled with the approach name (not blinded) - see docs/decisions.md D6.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import asdict
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..conflicts.compare import Conflict
from .schema import DecisionRow, Note, Vocabulary
from .then_text import then_sentence

VERDICT_COLUMNS = ["ra_verdict", "ra_comment", "doctor_verdict", "doctor_comment"]
NOT_RELEVANT = "–"
CONFLICT_HEADERS = {"situation": "patient situation", "guideline_1": "guideline 1", "says_1": "guideline 1 says",
                    "guideline_2": "guideline 2", "says_2": "guideline 2 says", "conflict_type": "conflict type",
                    "note": "plain-words note"}
# group -> (header colour, cell colour)
COLOURS = {"": ("FFFFFF", "FFFFFF"), "IF": ("9BC2E6", "DDEBF7"), "THEN": ("A9D08E", "E2EFDA"),
           "SOURCE": ("FFD966", "FFF2CC"), "AUTO CHECKS": ("BFBFBF", "EDEDED"), "REVIEW": ("F4B084", "FCE4D6")}
WIDTH = {"row_id": 20, "guideline": 14, "other conditions": 30, "THEN": 46, "drug class": 18, "specific drug": 16,
         "CHECKS": 34, "CAUTIONS": 34, "page": 22, "heading": 24, "quote": 60, "inferred because": 34,
         "qualifier check": 30, "IF check": 22, "THEN check": 30, "ra_comment": 26, "doctor_comment": 26}
EXPLAIN = {
    "THEN": "The action as one sentence. AND = give together; OR = choose one; 'one of (...)' = pick one drug.",
    "action": "start / add / switch / continue / stop / avoid (= do not start).",
    "specific drug": "Filled when the source names a drug, not the whole class.",
    "strength": "preferred / recommended / consider, from the source's wording.",
    "stated / inferred": "stated = every condition and the action are written in the quote.",
    "inferred because": "Why the model or the computer marked the row inferred.",
    "quote check": "100 = the quote is on the cited page word for word; lower = differences.",
    "qualifier check": "A condition word in the quote (if, unless, on, ...) that the IF may have dropped.",
    "IF check": "Every rule needs at least one IF condition.",
    "THEN check": "AND/OR read from '+' and '/' in the source - please confirm against the quote.",
}


def _notes(notes: Sequence[Note]) -> str:
    return "\n".join(f"[{n.basis}] {n.text}" for n in notes)


def _input_header(name: str, vocab: Vocabulary) -> str:
    spec = vocab.inputs[name]
    return f"{spec.label} ({spec.unit})" if spec.unit else spec.label


def rule_columns(row: DecisionRow, vocab: Vocabulary, checks: Mapping) -> dict[str, dict[str, str]]:
    """One line of the table as {group: {column: value}}, in display order."""
    if_part = {}
    for name in vocab.inputs:
        value = row.conditions.get(name)
        shown = vocab.display(name, value) if value else NOT_RELEVANT
        if value and name in row.scope_inputs:
            shown += " (from guideline scope)"
        if_part[_input_header(name, vocab)] = shown
    if_part["other conditions"] = "\n".join(row.other_conditions) or NOT_RELEVANT
    sentence, then_note = then_sentence(row)
    quote_rules = checks.get("quote_rules", [])
    score = round(checks.get("quote_score", 0))
    return {
        "": {"row_id": row.row_id, "guideline": row.guideline},
        "IF": if_part,
        "THEN": {
            "THEN": sentence,
            "action": row.action_type,
            "drug class": " / ".join(row.actions),
            "specific drug": ", ".join(row.specific_drugs) or "– (whole class)",
            "strength": row.strength,
            "CHECKS": _notes(row.checks),
            "CAUTIONS": _notes(row.cautions),
        },
        "SOURCE": {
            "page": "\n".join(f"{s.doc_id} p.{s.page}" for s in row.sources),
            "heading": "\n".join(s.heading for s in row.sources if s.heading),
            "quote": "\n".join(s.quote for s in row.sources),
        },
        "AUTO CHECKS": {
            "stated / inferred": row.label,
            "inferred because": "\n".join(row.label_reasons),
            "quote check": f"{score}" + (" - " + "; ".join(quote_rules) if quote_rules else ""),
            "qualifier check": "\n".join(checks.get("qualifiers", [])) or "OK",
            "IF check": "\n".join(checks.get("if_check", [])) or "OK",
            "THEN check": then_note or "OK",
        },
        "REVIEW": {c: "" for c in VERDICT_COLUMNS},
    }


def write_rule_sheet(wb, title: str, rows: Sequence[DecisionRow], vocab: Vocabulary,
                     row_checks: Mapping[str, Mapping]) -> None:
    """Rows in the WHO-style layout: group header row, column names, then one line per rule."""
    ws = wb.create_sheet(title)
    template = rule_columns(rows[0] if rows else DecisionRow("", {}, (), ()), vocab, {})
    thin = Side(style="thin", color="BFBFBF")
    spans, col = {}, 1
    for group, cols in template.items():
        header = COLOURS[group][0]
        ws.cell(row=1, column=col, value=group or None)
        if len(cols) > 1:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
        for i, name in enumerate(cols):
            top, head = ws.cell(row=1, column=col + i), ws.cell(row=2, column=col + i, value=name)
            top.fill = head.fill = PatternFill("solid", fgColor=header)
            top.font, head.font = Font(bold=True, size=12), Font(bold=True)
            top.alignment = Alignment(horizontal="center", vertical="center")
            head.alignment = Alignment(wrap_text=True, vertical="top")
            head.border = Border(bottom=thin)
        spans[group] = (col, list(cols))
        col += len(cols)

    for r_index, row in enumerate(rows, start=3):
        values = rule_columns(row, vocab, row_checks.get(row.row_id, {}))
        for group, (start, cols) in spans.items():
            fill = PatternFill("solid", fgColor=COLOURS[group][1])
            for i, name in enumerate(cols):
                cell = ws.cell(row=r_index, column=start + i, value=values[group][name])
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.fill = fill
                cell.border = Border(bottom=thin)

    for group, (start, cols) in spans.items():
        for i, name in enumerate(cols):
            letter = get_column_letter(start + i)
            ws.column_dimensions[letter].width = WIDTH.get(name, 14)
            # Hide patient inputs that no row in this sheet uses.
            if group == "IF" and rows and all(ws.cell(row=r, column=start + i).value == NOT_RELEVANT
                                              for r in range(3, len(rows) + 3)):
                ws.column_dimensions[letter].hidden = True
    ws.freeze_panes = "C3"  # both header rows, row_id and guideline stay visible
    ws.row_dimensions[2].height = 45


def write_how_to_read(wb, example: DecisionRow | None, vocab: Vocabulary, checks: Mapping) -> None:
    ws = wb.create_sheet("How to read", 0)
    bold = Font(bold=True)
    intro = [
        ("How to read this decision table", ""),
        ("", "Each line of the Rows sheet is one recommendation from one guideline: IF the patient matches every IF "
             "column, THEN the action applies. The column groups are coloured:"),
        ("IF", "Patient details the rule depends on. '–' = does not matter for this rule. Columns no rule uses are "
               "hidden. '(from guideline scope)' = true for every patient that guideline covers (e.g. KDIGO: CKD)."),
        ("THEN", "What to do. AND = give together (a combination); OR = choose one (alternatives); 'one of (...)' "
                 "= pick one drug from the list. Also drug class, a specific drug if the source names one, strength "
                 "(preferred / recommended / consider) and the CHECKS / CAUTIONS the source states or implies."),
        ("SOURCE", "Where it comes from: PDF and page, the table or list heading, and the exact quote."),
        ("AUTO CHECKS", "Computer checks, not clinical judgement - they point at rows worth a closer look."),
        ("REVIEW", "For you. Verdict: correct / incorrect / unsupported (the source does not say this). "
                   "Comment: what is wrong or missing."),
    ]
    r = 1
    for label, text in intro:
        a, b = ws.cell(row=r, column=1, value=label or None), ws.cell(row=r, column=2, value=text or None)
        if label in COLOURS:
            a.fill = PatternFill("solid", fgColor=COLOURS[label][0])
            a.font = bold
        r += 1
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)

    if example is not None:
        values = rule_columns(example, vocab, checks)
        r += 1
        ws.cell(row=r, column=1, value=f"Worked example: {example.row_id}").font = Font(bold=True, size=12)
        r += 1
        for c, head in enumerate(["Column", "Value in this row", "What it means"], start=1):
            ws.cell(row=r, column=c, value=head).font = bold
        r += 1
        for group, cols in values.items():
            if group == "REVIEW":
                continue
            for name, value in cols.items():
                if group == "IF" and value == NOT_RELEVANT:
                    continue  # only the inputs this rule uses
                ws.cell(row=r, column=1, value=f"{group}: {name}" if group else name)
                ws.cell(row=r, column=2, value=value)
                ws.cell(row=r, column=3, value=EXPLAIN.get(name, ""))
                if group:
                    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=COLOURS[group][1])
                r += 1
        used = [f"{k} = {v}" for k, v in values["IF"].items() if v != NOT_RELEVANT]
        ws.cell(row=r + 1, column=1, value="In words").font = bold
        ws.cell(row=r + 1, column=2, value=f"IF {'; '.join(used) or 'any patient'}, THEN {values['THEN']['THEN']} "
                                           f"({values['THEN']['strength']}) - {example.guideline}, "
                                           f"{values['SOURCE']['page']}.").font = Font(italic=True)
    for letter, width in (("A", 30), ("B", 75), ("C", 70)):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _conflicts_df(items: Sequence[Conflict], id_header: str) -> pd.DataFrame:
    df = pd.DataFrame([asdict(c) for c in items], columns=["conflict_id", *CONFLICT_HEADERS])
    return df.rename(columns={"conflict_id": id_header, **CONFLICT_HEADERS})


def _tidy_plain_sheet(ws, widths: Mapping[str, int]) -> None:
    """Bold, frozen header row and wrapped text on the sheets written by pandas."""
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, ws.max_column + 1):
        header = str(ws.cell(row=1, column=i).value or "")
        ws.column_dimensions[get_column_letter(i)].width = widths.get(header, 18 if len(header) < 20 else 30)
    ws.freeze_panes = "A2"


def export_workbook(
    path: str | Path,
    rows: Sequence[DecisionRow],
    vocab: Vocabulary,
    row_checks: Mapping[str, Mapping],
    coverage: Sequence[Mapping],
    conflicts: Sequence[Conflict],
    compare_blocks: Sequence[Mapping],
    skipped: Sequence[Mapping],
    dose_rows: Sequence[DecisionRow] = (),
    example_row_id: str | None = None,
) -> Path:
    path = Path(path)
    conflicts_df = _conflicts_df(conflicts, "conflict_id")
    conflicts_df["doctor_decision"] = ""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        wb = writer.book
        example = next((r for r in rows if r.row_id == example_row_id), rows[0] if rows else None)
        write_how_to_read(wb, example, vocab, row_checks.get(example.row_id, {}) if example else {})
        write_rule_sheet(wb, "Rows", rows, vocab, row_checks)
        pd.DataFrame(coverage).to_excel(writer, sheet_name="Coverage", index=False)
        conflicts_df.to_excel(writer, sheet_name="Conflicts", index=False)
        pd.DataFrame(compare_blocks).to_excel(writer, sheet_name="Compare by drug", index=False)
        write_rule_sheet(wb, "For dose table", dose_rows, vocab, row_checks)
        pd.DataFrame(skipped, columns=["guideline", "category", "doc_id", "page", "text", "possible_recommendation"]
                     ).to_excel(writer, sheet_name="Skipped", index=False)
        widths = {"guideline 1 says": 60, "guideline 2 says": 60, "plain-words note": 50, "text": 60,
                  **{g: 55 for g in (compare_blocks[0] if compare_blocks else {}) if g not in ("drug class", "decision")}}
        for name in ("Coverage", "Conflicts", "Compare by drug", "Skipped"):
            _tidy_plain_sheet(wb[name], widths)
    return path


def read_review_files(out_dir: str | Path, who: str = "ra") -> dict[str, list[str]]:
    """Return {approach: [verdict, ...]} from the Rows sheets, for who = "ra" or "doctor".

    Row 1 of the Rows sheet is the group header (IF / THEN / ...); the column names are row 2.
    """
    verdicts = {}
    for path in sorted(Path(out_dir).glob("Approach_*.xlsx")):
        sheet = pd.read_excel(path, sheet_name="Rows", header=1, dtype=str).fillna("")
        given = sheet[f"{who}_verdict"].str.strip()
        verdicts[path.stem.removeprefix("Approach_")] = given[given != ""].tolist()
    return verdicts
